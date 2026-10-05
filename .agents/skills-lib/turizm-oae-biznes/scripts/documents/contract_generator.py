#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
Генератор договоров с автозаполнением.

Поддерживаемые типы договоров:
1. Договор на тур/экскурсию
2. Агентский договор
3. Договор аренды авто
4. Договор на яхту

Функции:
- Шаблоны в templates/contracts/
- Автозаполнение из данных клиента
- Генерация PDF (reportlab)
- Генерация DOCX (python-docx)
- Цифровая подпись (опционально)
- QR код с ID договора
"""

import json
import os
import re
import hashlib
import qrcode
from datetime import datetime, timedelta
from pathlib import Path
from typing import Dict, List, Optional, Any, Union
from dataclasses import dataclass, field, asdict
from enum import Enum
import uuid

# Попытка импорта библиотек для генерации документов
try:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm, cm
    from reportlab.platypus import (
        SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle,
        Image, PageBreak, KeepTogether
    )
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    REPORTLAB_AVAILABLE = True
except ImportError:
    REPORTLAB_AVAILABLE = False

try:
    from docx import Document
    from docx.shared import Inches, Pt, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

try:
    from cryptography.hazmat.primitives import hashes, serialization
    from cryptography.hazmat.primitives.asymmetric import rsa, padding
    from cryptography.hazmat.backends import default_backend
    CRYPTO_AVAILABLE = True
except ImportError:
    CRYPTO_AVAILABLE = False

# Импорт конфигурации
try:
    from config import CHATS_DIR, ensure_directories
except ImportError:
    CHATS_DIR = Path("D:/Downloads/Chats")
    def ensure_directories():
        pass

# ═══════════════════════════════════════════════════════════════
# КОНСТАНТЫ И ПУТИ
# ═══════════════════════════════════════════════════════════════

SCRIPTS_DIR = Path(__file__).parent
TEMPLATES_DIR = SCRIPTS_DIR / "templates" / "contracts"
CONTRACTS_DIR = CHATS_DIR / "_договоры"
REGISTRY_FILE = CONTRACTS_DIR / "contract_registry.json"

# Компания по умолчанию
DEFAULT_COMPANY = {
    "name_ru": "Marsel Luxury Car Rental",
    "name_en": "Marsel Luxury Car Rental LLC",
    "legal_address": "Dubai, UAE, Business Bay, Churchill Tower",
    "postal_address": "Dubai, UAE, Business Bay, Churchill Tower",
    "phone": "+971 50 770 5321",
    "email": "info@marsel-luxury.ae",
    "website": "www.marsel-luxury.ae",
    "trade_license": "1234567",
    "vat_number": "100123456700003",
    "bank_name": "Emirates NBD",
    "iban": "AE123456789012345678901",
    "swift": "EABORWRXXXX",
    "director": "Марсель Шакиров",
    "director_en": "Marsel Shakirov"
}


class ContractType(Enum):
    """Типы договоров."""
    TOUR = "tour"              # Договор на тур/экскурсию
    AGENT = "agent"            # Агентский договор
    CAR_RENTAL = "car_rental"  # Договор аренды авто
    YACHT = "yacht"            # Договор на яхту


class ContractStatus(Enum):
    """Статусы договора."""
    DRAFT = "draft"            # Черновик
    PENDING = "pending"        # Ожидает подписания
    SIGNED = "signed"          # Подписан
    ACTIVE = "active"          # Действует
    COMPLETED = "completed"    # Исполнен
    CANCELLED = "cancelled"    # Отменён


@dataclass
class ClientData:
    """Данные клиента/контрагента."""
    full_name: str
    full_name_en: Optional[str] = None
    passport_number: Optional[str] = None
    passport_issued: Optional[str] = None
    passport_expiry: Optional[str] = None
    nationality: Optional[str] = None
    date_of_birth: Optional[str] = None
    phone: Optional[str] = None
    email: Optional[str] = None
    address: Optional[str] = None

    # Для юридических лиц
    company_name: Optional[str] = None
    company_name_en: Optional[str] = None
    trade_license: Optional[str] = None
    vat_number: Optional[str] = None
    legal_address: Optional[str] = None
    bank_name: Optional[str] = None
    iban: Optional[str] = None
    swift: Optional[str] = None
    authorized_person: Optional[str] = None
    position: Optional[str] = None


@dataclass
class TourServiceData:
    """Данные услуги тура/экскурсии."""
    name: str
    name_en: Optional[str] = None
    description: Optional[str] = None
    date: str = ""
    time: str = ""
    duration: str = ""
    pickup_location: Optional[str] = None
    dropoff_location: Optional[str] = None
    participants: int = 1
    children: int = 0
    price_per_person: float = 0
    total_price: float = 0
    currency: str = "AED"
    includes: List[str] = field(default_factory=list)
    excludes: List[str] = field(default_factory=list)
    notes: Optional[str] = None


@dataclass
class CarRentalData:
    """Данные аренды автомобиля."""
    vehicle_make: str
    vehicle_model: str
    vehicle_year: int
    plate_number: str
    vin_number: Optional[str] = None
    color: Optional[str] = None
    mileage_start: int = 0
    mileage_limit: Optional[int] = None

    rental_start: str = ""
    rental_end: str = ""
    pickup_location: str = ""
    dropoff_location: str = ""

    daily_rate: float = 0
    weekly_rate: Optional[float] = None
    monthly_rate: Optional[float] = None
    total_days: int = 1
    total_price: float = 0
    currency: str = "AED"

    deposit_amount: float = 0
    deposit_type: str = "cash"  # cash, card_hold, bank_transfer

    insurance_type: str = "basic"  # basic, full, premium
    insurance_included: bool = True
    excess_amount: float = 0

    additional_driver: bool = False
    additional_driver_fee: float = 0

    fuel_policy: str = "full_to_full"  # full_to_full, same_to_same

    special_conditions: List[str] = field(default_factory=list)


@dataclass
class YachtCharterData:
    """Данные чартера яхты."""
    yacht_name: str
    yacht_type: str  # motor, sailing, catamaran
    yacht_length: str
    capacity: int

    charter_date: str = ""
    departure_time: str = ""
    return_time: str = ""
    duration_hours: int = 4

    departure_marina: str = ""
    route_description: Optional[str] = None
    stops: List[str] = field(default_factory=list)

    guests: int = 1
    crew_included: bool = True
    captain_name: Optional[str] = None

    base_price: float = 0
    per_hour_rate: float = 0
    total_price: float = 0
    currency: str = "AED"

    deposit_amount: float = 0

    catering_included: bool = False
    catering_description: Optional[str] = None
    catering_price: float = 0

    fuel_included: bool = True
    fuel_policy: Optional[str] = None

    special_requests: List[str] = field(default_factory=list)


@dataclass
class AgentContractData:
    """Данные агентского договора."""
    agency_name: str
    agency_name_en: Optional[str] = None
    trade_license: str = ""
    vat_number: Optional[str] = None
    legal_address: str = ""

    contact_person: str = ""
    contact_position: str = ""
    contact_phone: str = ""
    contact_email: str = ""

    commission_rate: float = 10.0  # Процент комиссии
    commission_type: str = "percentage"  # percentage, fixed, tiered

    payment_terms: str = "prepaid"  # prepaid, postpaid, mixed
    payment_period: int = 7  # Дней на оплату

    contract_start: str = ""
    contract_end: str = ""
    auto_renewal: bool = True
    notice_period: int = 30  # Дней для уведомления о расторжении

    services_covered: List[str] = field(default_factory=list)
    exclusive: bool = False
    territory: Optional[str] = None

    bank_name: str = ""
    iban: str = ""
    swift: str = ""


@dataclass
class ContractData:
    """Основные данные договора."""
    contract_id: str
    contract_type: ContractType
    contract_number: str

    created_at: str = ""
    signed_at: Optional[str] = None
    valid_from: str = ""
    valid_until: str = ""

    status: ContractStatus = ContractStatus.DRAFT

    client: Optional[ClientData] = None
    company: Dict[str, str] = field(default_factory=lambda: DEFAULT_COMPANY.copy())

    # Специфичные данные для типа договора
    tour_data: Optional[TourServiceData] = None
    car_rental_data: Optional[CarRentalData] = None
    yacht_data: Optional[YachtCharterData] = None
    agent_data: Optional[AgentContractData] = None

    # Финансы
    total_amount: float = 0
    currency: str = "AED"
    prepayment_amount: float = 0
    prepayment_percent: float = 50

    # Условия отмены
    cancellation_policy: str = "standard"
    cancellation_terms: List[Dict[str, Any]] = field(default_factory=list)

    # Подпись
    digital_signature: Optional[str] = None
    signature_date: Optional[str] = None

    # Дополнительно
    language: str = "ru"  # ru, en, both
    notes: Optional[str] = None
    attachments: List[str] = field(default_factory=list)


# ═══════════════════════════════════════════════════════════════
# ГЕНЕРАТОР НОМЕРОВ ДОГОВОРОВ
# ═══════════════════════════════════════════════════════════════

class ContractNumberGenerator:
    """Генератор уникальных номеров договоров."""

    PREFIXES = {
        ContractType.TOUR: "TUR",
        ContractType.AGENT: "AGT",
        ContractType.CAR_RENTAL: "CAR",
        ContractType.YACHT: "YCH"
    }

    def __init__(self, registry_file: Path = REGISTRY_FILE):
        self.registry_file = registry_file
        self.counters = self._load_counters()

    def _load_counters(self) -> Dict[str, int]:
        """Загрузить счётчики из реестра."""
        if self.registry_file.exists():
            try:
                with open(self.registry_file, 'r', encoding='utf-8') as f:
                    data = json.load(f)
                    return data.get("counters", {})
            except:
                pass
        return {}

    def _save_counters(self):
        """Сохранить счётчики в реестр."""
        registry = self._load_registry()
        registry["counters"] = self.counters
        self._save_registry(registry)

    def _load_registry(self) -> Dict:
        """Загрузить реестр."""
        if self.registry_file.exists():
            try:
                with open(self.registry_file, 'r', encoding='utf-8') as f:
                    return json.load(f)
            except:
                pass
        return {"contracts": [], "counters": {}}

    def _save_registry(self, registry: Dict):
        """Сохранить реестр."""
        self.registry_file.parent.mkdir(parents=True, exist_ok=True)
        with open(self.registry_file, 'w', encoding='utf-8') as f:
            json.dump(registry, f, ensure_ascii=False, indent=2)

    def generate(self, contract_type: ContractType) -> str:
        """Генерировать номер договора."""
        prefix = self.PREFIXES.get(contract_type, "CTR")
        year = datetime.now().strftime("%y")
        month = datetime.now().strftime("%m")

        key = f"{prefix}-{year}"
        current = self.counters.get(key, 0) + 1
        self.counters[key] = current
        self._save_counters()

        return f"{prefix}-{year}{month}-{current:04d}"

    def generate_id(self) -> str:
        """Генерировать уникальный ID договора."""
        return str(uuid.uuid4())[:12].upper()


# ═══════════════════════════════════════════════════════════════
# РЕЕСТР ДОГОВОРОВ
# ═══════════════════════════════════════════════════════════════

class ContractRegistry:
    """Реестр всех договоров."""

    def __init__(self, registry_file: Path = REGISTRY_FILE):
        self.registry_file = registry_file
        self._ensure_file()

    def _ensure_file(self):
        """Убедиться что файл реестра существует."""
        if not self.registry_file.exists():
            self.registry_file.parent.mkdir(parents=True, exist_ok=True)
            self._save({"contracts": [], "counters": {}})

    def _load(self) -> Dict:
        """Загрузить реестр."""
        try:
            with open(self.registry_file, 'r', encoding='utf-8') as f:
                return json.load(f)
        except:
            return {"contracts": [], "counters": {}}

    def _save(self, data: Dict):
        """Сохранить реестр."""
        with open(self.registry_file, 'w', encoding='utf-8') as f:
            json.dump(data, f, ensure_ascii=False, indent=2)

    def add_contract(self, contract: ContractData) -> bool:
        """Добавить договор в реестр."""
        registry = self._load()

        # Проверка на дубликат
        for c in registry["contracts"]:
            if c["contract_id"] == contract.contract_id:
                return False

        # Преобразование в словарь
        contract_dict = self._contract_to_dict(contract)
        registry["contracts"].append(contract_dict)

        self._save(registry)
        return True

    def update_contract(self, contract: ContractData) -> bool:
        """Обновить договор в реестре."""
        registry = self._load()

        for i, c in enumerate(registry["contracts"]):
            if c["contract_id"] == contract.contract_id:
                registry["contracts"][i] = self._contract_to_dict(contract)
                self._save(registry)
                return True

        return False

    def get_contract(self, contract_id: str) -> Optional[Dict]:
        """Получить договор по ID."""
        registry = self._load()
        for c in registry["contracts"]:
            if c["contract_id"] == contract_id:
                return c
        return None

    def find_contracts(
        self,
        contract_type: Optional[ContractType] = None,
        status: Optional[ContractStatus] = None,
        client_name: Optional[str] = None,
        date_from: Optional[str] = None,
        date_to: Optional[str] = None
    ) -> List[Dict]:
        """Поиск договоров по критериям."""
        registry = self._load()
        results = []

        for c in registry["contracts"]:
            if contract_type and c.get("contract_type") != contract_type.value:
                continue
            if status and c.get("status") != status.value:
                continue
            if client_name:
                client = c.get("client", {})
                if client_name.lower() not in client.get("full_name", "").lower():
                    continue
            if date_from and c.get("created_at", "") < date_from:
                continue
            if date_to and c.get("created_at", "") > date_to:
                continue

            results.append(c)

        return results

    def get_statistics(self) -> Dict:
        """Получить статистику по договорам."""
        registry = self._load()
        contracts = registry["contracts"]

        stats = {
            "total": len(contracts),
            "by_type": {},
            "by_status": {},
            "total_amount": 0,
            "this_month": 0,
            "this_year": 0
        }

        current_month = datetime.now().strftime("%Y-%m")
        current_year = datetime.now().strftime("%Y")

        for c in contracts:
            # По типу
            ctype = c.get("contract_type", "unknown")
            stats["by_type"][ctype] = stats["by_type"].get(ctype, 0) + 1

            # По статусу
            status = c.get("status", "unknown")
            stats["by_status"][status] = stats["by_status"].get(status, 0) + 1

            # Сумма
            stats["total_amount"] += c.get("total_amount", 0)

            # За период
            created = c.get("created_at", "")
            if created.startswith(current_month):
                stats["this_month"] += 1
            if created.startswith(current_year):
                stats["this_year"] += 1

        return stats

    def _contract_to_dict(self, contract: ContractData) -> Dict:
        """Преобразовать ContractData в словарь."""
        result = {
            "contract_id": contract.contract_id,
            "contract_type": contract.contract_type.value,
            "contract_number": contract.contract_number,
            "created_at": contract.created_at,
            "signed_at": contract.signed_at,
            "valid_from": contract.valid_from,
            "valid_until": contract.valid_until,
            "status": contract.status.value,
            "total_amount": contract.total_amount,
            "currency": contract.currency,
            "language": contract.language,
            "notes": contract.notes
        }

        if contract.client:
            result["client"] = asdict(contract.client)
        if contract.company:
            result["company"] = contract.company

        return result


# ═══════════════════════════════════════════════════════════════
# ГЕНЕРАТОР QR КОДОВ
# ═══════════════════════════════════════════════════════════════

class QRCodeGenerator:
    """Генератор QR кодов для договоров."""

    @staticmethod
    def generate(
        contract_id: str,
        contract_number: str,
        output_path: Path,
        size: int = 200
    ) -> Path:
        """Генерировать QR код с данными договора."""
        # Данные для QR
        data = json.dumps({
            "id": contract_id,
            "number": contract_number,
            "verify": f"https://verify.marsel-luxury.ae/contract/{contract_id}"
        }, ensure_ascii=False)

        # Создание QR
        qr = qrcode.QRCode(
            version=1,
            error_correction=qrcode.constants.ERROR_CORRECT_M,
            box_size=10,
            border=2
        )
        qr.add_data(data)
        qr.make(fit=True)

        img = qr.make_image(fill_color="black", back_color="white")

        # Сохранение
        qr_path = output_path / f"qr_{contract_id}.png"
        img.save(qr_path)

        return qr_path


# ═══════════════════════════════════════════════════════════════
# ЦИФРОВАЯ ПОДПИСЬ (ОПЦИОНАЛЬНО)
# ═══════════════════════════════════════════════════════════════

class DigitalSigner:
    """Цифровая подпись документов."""

    def __init__(self, private_key_path: Optional[Path] = None):
        self.private_key = None
        self.public_key = None

        if private_key_path and CRYPTO_AVAILABLE:
            self._load_keys(private_key_path)

    def _load_keys(self, private_key_path: Path):
        """Загрузить ключи из файла."""
        if not private_key_path.exists():
            return

        with open(private_key_path, 'rb') as f:
            self.private_key = serialization.load_pem_private_key(
                f.read(),
                password=None,
                backend=default_backend()
            )
            self.public_key = self.private_key.public_key()

    def generate_keys(self, output_dir: Path) -> tuple:
        """Генерировать новую пару ключей."""
        if not CRYPTO_AVAILABLE:
            raise ImportError("cryptography library not installed")

        private_key = rsa.generate_private_key(
            public_exponent=65537,
            key_size=2048,
            backend=default_backend()
        )
        public_key = private_key.public_key()

        # Сохранение приватного ключа
        private_path = output_dir / "contract_private_key.pem"
        with open(private_path, 'wb') as f:
            f.write(private_key.private_bytes(
                encoding=serialization.Encoding.PEM,
                format=serialization.PrivateFormat.PKCS8,
                encryption_algorithm=serialization.NoEncryption()
            ))

        # Сохранение публичного ключа
        public_path = output_dir / "contract_public_key.pem"
        with open(public_path, 'wb') as f:
            f.write(public_key.public_bytes(
                encoding=serialization.Encoding.PEM,
                format=serialization.PublicFormat.SubjectPublicKeyInfo
            ))

        self.private_key = private_key
        self.public_key = public_key

        return private_path, public_path

    def sign_document(self, document_hash: bytes) -> Optional[bytes]:
        """Подписать хеш документа."""
        if not self.private_key or not CRYPTO_AVAILABLE:
            return None

        signature = self.private_key.sign(
            document_hash,
            padding.PSS(
                mgf=padding.MGF1(hashes.SHA256()),
                salt_length=padding.PSS.MAX_LENGTH
            ),
            hashes.SHA256()
        )
        return signature

    def verify_signature(self, document_hash: bytes, signature: bytes) -> bool:
        """Проверить подпись."""
        if not self.public_key or not CRYPTO_AVAILABLE:
            return False

        try:
            self.public_key.verify(
                signature,
                document_hash,
                padding.PSS(
                    mgf=padding.MGF1(hashes.SHA256()),
                    salt_length=padding.PSS.MAX_LENGTH
                ),
                hashes.SHA256()
            )
            return True
        except:
            return False

    @staticmethod
    def hash_file(file_path: Path) -> bytes:
        """Вычислить хеш файла."""
        sha256 = hashlib.sha256()
        with open(file_path, 'rb') as f:
            for chunk in iter(lambda: f.read(4096), b''):
                sha256.update(chunk)
        return sha256.digest()


# ═══════════════════════════════════════════════════════════════
# ШАБЛОНЫ ДОГОВОРОВ
# ═══════════════════════════════════════════════════════════════

class ContractTemplates:
    """Шаблоны текстов договоров."""

    # Стандартные условия отмены
    CANCELLATION_POLICIES = {
        "strict": [
            {"days_before": 7, "refund_percent": 0},
            {"days_before": 14, "refund_percent": 50},
            {"days_before": 30, "refund_percent": 100}
        ],
        "standard": [
            {"days_before": 3, "refund_percent": 0},
            {"days_before": 7, "refund_percent": 50},
            {"days_before": 14, "refund_percent": 100}
        ],
        "flexible": [
            {"days_before": 1, "refund_percent": 50},
            {"days_before": 3, "refund_percent": 100}
        ],
        "non_refundable": [
            {"days_before": 0, "refund_percent": 0}
        ]
    }

    @staticmethod
    def get_tour_contract_text(contract: ContractData, language: str = "ru") -> str:
        """Получить текст договора на тур."""
        if language == "ru":
            return ContractTemplates._tour_contract_ru(contract)
        else:
            return ContractTemplates._tour_contract_en(contract)

    @staticmethod
    def _tour_contract_ru(contract: ContractData) -> str:
        """Текст договора на тур (русский)."""
        tour = contract.tour_data
        client = contract.client
        company = contract.company

        text = f"""
ДОГОВОР НА ОКАЗАНИЕ ТУРИСТИЧЕСКИХ УСЛУГ
№ {contract.contract_number}

г. Дубай                                                     {contract.created_at}

{company['name_ru']}, именуемое в дальнейшем «Исполнитель», в лице директора
{company['director']}, действующего на основании Устава, с одной стороны, и

{client.full_name}, паспорт {client.passport_number or '________________'},
именуемый(ая) в дальнейшем «Заказчик», с другой стороны,

заключили настоящий Договор о нижеследующем:

1. ПРЕДМЕТ ДОГОВОРА

1.1. Исполнитель обязуется оказать Заказчику следующие туристические услуги:

    Название: {tour.name}
    Дата проведения: {tour.date}
    Время начала: {tour.time}
    Продолжительность: {tour.duration}
    Количество участников: {tour.participants} взрослых, {tour.children} детей
    Место встречи: {tour.pickup_location or 'Согласуется дополнительно'}
    Место окончания: {tour.dropoff_location or 'Место встречи'}

1.2. В стоимость услуг включено:
{chr(10).join('    - ' + item for item in tour.includes) if tour.includes else '    - Согласно описанию тура'}

1.3. В стоимость услуг НЕ включено:
{chr(10).join('    - ' + item for item in tour.excludes) if tour.excludes else '    - Личные расходы'}

2. СТОИМОСТЬ УСЛУГ И ПОРЯДОК ОПЛАТЫ

2.1. Стоимость услуг составляет: {tour.total_price:,.2f} {tour.currency}
     (Цена за 1 человека: {tour.price_per_person:,.2f} {tour.currency})

2.2. Порядок оплаты:
     - Предоплата {contract.prepayment_percent}%: {contract.prepayment_amount:,.2f} {contract.currency}
     - Остаток: {contract.total_amount - contract.prepayment_amount:,.2f} {contract.currency}

2.3. Оплата производится:
     - Наличными в {contract.currency}
     - Банковским переводом на реквизиты Исполнителя
     - Банковской картой (+ 3% комиссия)

3. ПРАВА И ОБЯЗАННОСТИ СТОРОН

3.1. Исполнитель обязуется:
     - Оказать услуги качественно и в полном объёме
     - Обеспечить безопасность участников
     - Предоставить квалифицированного гида/водителя
     - Своевременно информировать об изменениях

3.2. Заказчик обязуется:
     - Своевременно произвести оплату
     - Прибыть в указанное место и время
     - Соблюдать правила безопасности
     - Иметь при себе документ, удостоверяющий личность

4. УСЛОВИЯ ОТМЕНЫ И ВОЗВРАТА

{ContractTemplates._format_cancellation_policy(contract.cancellation_terms or ContractTemplates.CANCELLATION_POLICIES['standard'])}

5. ОТВЕТСТВЕННОСТЬ СТОРОН

5.1. За неисполнение или ненадлежащее исполнение обязательств стороны несут
     ответственность в соответствии с законодательством ОАЭ.

5.2. Исполнитель не несёт ответственности за:
     - Действия третьих лиц
     - Форс-мажорные обстоятельства
     - Последствия несоблюдения Заказчиком правил безопасности

6. ПРОЧИЕ УСЛОВИЯ

6.1. Настоящий Договор вступает в силу с момента подписания.
6.2. Все споры разрешаются путём переговоров.
6.3. Договор составлен в двух экземплярах.

7. РЕКВИЗИТЫ И ПОДПИСИ СТОРОН

ИСПОЛНИТЕЛЬ:                              ЗАКАЗЧИК:
{company['name_ru']}                      {client.full_name}
{company['legal_address']}                {client.address or ''}
Тел: {company['phone']}                   Тел: {client.phone or ''}
Email: {company['email']}                 Email: {client.email or ''}
Trade License: {company['trade_license']}
IBAN: {company['iban']}

_______________________                   _______________________
{company['director']}
Директор

Дата: _______________                     Дата: _______________
"""
        return text

    @staticmethod
    def _tour_contract_en(contract: ContractData) -> str:
        """Текст договора на тур (английский)."""
        tour = contract.tour_data
        client = contract.client
        company = contract.company

        text = f"""
TOURISM SERVICES AGREEMENT
No. {contract.contract_number}

Dubai                                                        {contract.created_at}

{company['name_en']}, hereinafter referred to as the "Service Provider",
represented by Director {company['director_en']}, acting on the basis of the Charter,
on the one hand, and

{client.full_name_en or client.full_name}, passport {client.passport_number or '________________'},
hereinafter referred to as the "Client", on the other hand,

have concluded this Agreement as follows:

1. SUBJECT OF THE AGREEMENT

1.1. The Service Provider undertakes to provide the following tourism services:

    Service Name: {tour.name_en or tour.name}
    Date: {tour.date}
    Time: {tour.time}
    Duration: {tour.duration}
    Participants: {tour.participants} adults, {tour.children} children
    Pick-up Location: {tour.pickup_location or 'To be confirmed'}
    Drop-off Location: {tour.dropoff_location or 'Pick-up location'}

1.2. The price includes:
{chr(10).join('    - ' + item for item in tour.includes) if tour.includes else '    - As per tour description'}

1.3. The price does NOT include:
{chr(10).join('    - ' + item for item in tour.excludes) if tour.excludes else '    - Personal expenses'}

2. PRICE AND PAYMENT TERMS

2.1. Total Price: {tour.total_price:,.2f} {tour.currency}
     (Price per person: {tour.price_per_person:,.2f} {tour.currency})

2.2. Payment Schedule:
     - Advance payment {contract.prepayment_percent}%: {contract.prepayment_amount:,.2f} {contract.currency}
     - Balance: {contract.total_amount - contract.prepayment_amount:,.2f} {contract.currency}

2.3. Payment Methods:
     - Cash in {contract.currency}
     - Bank transfer to Service Provider's account
     - Credit/Debit card (+ 3% processing fee)

3. RIGHTS AND OBLIGATIONS

3.1. The Service Provider undertakes to:
     - Provide services professionally and in full
     - Ensure participants' safety
     - Provide qualified guide/driver
     - Inform about any changes in advance

3.2. The Client undertakes to:
     - Make timely payments
     - Arrive at the specified location and time
     - Follow safety instructions
     - Carry valid identification

4. CANCELLATION AND REFUND POLICY

{ContractTemplates._format_cancellation_policy_en(contract.cancellation_terms or ContractTemplates.CANCELLATION_POLICIES['standard'])}

5. LIABILITY

5.1. Both parties are liable for non-performance or improper performance of
     obligations in accordance with UAE law.

5.2. The Service Provider is not liable for:
     - Actions of third parties
     - Force majeure circumstances
     - Consequences of Client's non-compliance with safety rules

6. GENERAL PROVISIONS

6.1. This Agreement comes into force upon signing.
6.2. All disputes shall be resolved through negotiations.
6.3. This Agreement is made in two copies.

7. DETAILS AND SIGNATURES

SERVICE PROVIDER:                         CLIENT:
{company['name_en']}                      {client.full_name_en or client.full_name}
{company['legal_address']}                {client.address or ''}
Tel: {company['phone']}                   Tel: {client.phone or ''}
Email: {company['email']}                 Email: {client.email or ''}
Trade License: {company['trade_license']}
IBAN: {company['iban']}

_______________________                   _______________________
{company['director_en']}
Director

Date: _______________                     Date: _______________
"""
        return text

    @staticmethod
    def get_car_rental_contract_text(contract: ContractData, language: str = "ru") -> str:
        """Получить текст договора аренды авто."""
        if language == "ru":
            return ContractTemplates._car_rental_contract_ru(contract)
        else:
            return ContractTemplates._car_rental_contract_en(contract)

    @staticmethod
    def _car_rental_contract_ru(contract: ContractData) -> str:
        """Текст договора аренды авто (русский)."""
        car = contract.car_rental_data
        client = contract.client
        company = contract.company

        text = f"""
ДОГОВОР АРЕНДЫ ТРАНСПОРТНОГО СРЕДСТВА
№ {contract.contract_number}

г. Дубай                                                     {contract.created_at}

{company['name_ru']}, именуемое в дальнейшем «Арендодатель», в лице директора
{company['director']}, действующего на основании Устава, с одной стороны, и

{client.full_name}, паспорт {client.passport_number or '________________'},
водительское удостоверение ________________,
именуемый(ая) в дальнейшем «Арендатор», с другой стороны,

заключили настоящий Договор о нижеследующем:

1. ПРЕДМЕТ ДОГОВОРА

1.1. Арендодатель передаёт, а Арендатор принимает во временное владение и
     пользование следующее транспортное средство:

     Марка/Модель: {car.vehicle_make} {car.vehicle_model}
     Год выпуска: {car.vehicle_year}
     Цвет: {car.color or 'N/A'}
     Государственный номер: {car.plate_number}
     VIN: {car.vin_number or 'N/A'}
     Пробег на момент передачи: {car.mileage_start:,} км

2. СРОК АРЕНДЫ

2.1. Период аренды:
     Начало: {car.rental_start}
     Окончание: {car.rental_end}
     Количество дней: {car.total_days}

2.2. Место получения: {car.pickup_location}
2.3. Место возврата: {car.dropoff_location}

3. АРЕНДНАЯ ПЛАТА И ПОРЯДОК РАСЧЁТОВ

3.1. Стоимость аренды:
     - Суточная ставка: {car.daily_rate:,.2f} {car.currency}
     - Общая стоимость: {car.total_price:,.2f} {car.currency}

3.2. Залоговый депозит: {car.deposit_amount:,.2f} {car.currency}
     Тип депозита: {car.deposit_type}

3.3. Порядок оплаты:
     - Предоплата: {contract.prepayment_amount:,.2f} {contract.currency}
     - Остаток при получении авто

4. СТРАХОВАНИЕ

4.1. Тип страховки: {car.insurance_type}
4.2. Страховка включена: {'Да' if car.insurance_included else 'Нет'}
4.3. Франшиза: {car.excess_amount:,.2f} {car.currency}

5. УСЛОВИЯ ЭКСПЛУАТАЦИИ

5.1. Лимит пробега: {f'{car.mileage_limit:,} км' if car.mileage_limit else 'Без ограничений'}
5.2. Топливная политика: {car.fuel_policy}

5.3. ЗАПРЕЩАЕТСЯ:
     - Передавать управление третьим лицам (если не указан доп. водитель)
     - Использовать в качестве такси или для обучения вождению
     - Выезжать за пределы ОАЭ без письменного согласия
     - Курить в салоне автомобиля
     - Перевозить животных без специального разрешения
     - Использовать авто в нетрезвом состоянии

5.4. Дополнительный водитель: {'Да (+' + str(car.additional_driver_fee) + ' ' + car.currency + ')' if car.additional_driver else 'Нет'}

6. ОТВЕТСТВЕННОСТЬ АРЕНДАТОРА

6.1. Арендатор несёт ответственность за:
     - Повреждения автомобиля, не покрытые страховкой
     - Штрафы за нарушение ПДД
     - Утерю документов, ключей, номерных знаков
     - Ущерб интерьеру (прожоги, пятна, разрывы)

6.2. В случае ДТП Арендатор обязан:
     - Немедленно сообщить Арендодателю
     - Вызвать полицию и получить протокол
     - Не покидать место происшествия

7. ВОЗВРАТ АВТОМОБИЛЯ

7.1. Автомобиль должен быть возвращён:
     - В указанное время и место
     - В чистом состоянии
     - С тем же уровнем топлива
     - Со всеми документами и ключами

7.2. Просрочка возврата: штраф в размере суточной ставки за каждый день

{f'''8. ОСОБЫЕ УСЛОВИЯ

{chr(10).join('- ' + cond for cond in car.special_conditions)}
''' if car.special_conditions else ''}

9. РЕКВИЗИТЫ И ПОДПИСИ СТОРОН

АРЕНДОДАТЕЛЬ:                             АРЕНДАТОР:
{company['name_ru']}                      {client.full_name}
{company['legal_address']}                Паспорт: {client.passport_number or ''}
Тел: {company['phone']}                   Тел: {client.phone or ''}
Trade License: {company['trade_license']}
IBAN: {company['iban']}

_______________________                   _______________________
{company['director']}
Директор

Автомобиль передан: _______________       Автомобиль принят: _______________
Время: _______________                    Время: _______________

ВОЗВРАТ АВТОМОБИЛЯ

Дата/время возврата: _______________
Пробег при возврате: _______________ км
Уровень топлива: _______________
Состояние: _______________
Замечания: _______________

Подпись Арендодателя: _______________     Подпись Арендатора: _______________
"""
        return text

    @staticmethod
    def _car_rental_contract_en(contract: ContractData) -> str:
        """Текст договора аренды авто (английский)."""
        car = contract.car_rental_data
        client = contract.client
        company = contract.company

        text = f"""
VEHICLE RENTAL AGREEMENT
No. {contract.contract_number}

Dubai                                                        {contract.created_at}

{company['name_en']}, hereinafter referred to as the "Lessor",
represented by Director {company['director_en']}, on the one hand, and

{client.full_name_en or client.full_name}, passport {client.passport_number or '________________'},
driving license ________________,
hereinafter referred to as the "Lessee", on the other hand,

have concluded this Agreement as follows:

1. SUBJECT OF THE AGREEMENT

1.1. The Lessor transfers, and the Lessee accepts for temporary possession and
     use the following vehicle:

     Make/Model: {car.vehicle_make} {car.vehicle_model}
     Year: {car.vehicle_year}
     Color: {car.color or 'N/A'}
     License Plate: {car.plate_number}
     VIN: {car.vin_number or 'N/A'}
     Mileage at handover: {car.mileage_start:,} km

2. RENTAL PERIOD

2.1. Rental Period:
     Start: {car.rental_start}
     End: {car.rental_end}
     Total Days: {car.total_days}

2.2. Pick-up Location: {car.pickup_location}
2.3. Drop-off Location: {car.dropoff_location}

3. RENTAL CHARGES AND PAYMENT

3.1. Rental Rates:
     - Daily Rate: {car.daily_rate:,.2f} {car.currency}
     - Total Amount: {car.total_price:,.2f} {car.currency}

3.2. Security Deposit: {car.deposit_amount:,.2f} {car.currency}
     Deposit Type: {car.deposit_type}

3.3. Payment Terms:
     - Advance Payment: {contract.prepayment_amount:,.2f} {contract.currency}
     - Balance upon vehicle collection

4. INSURANCE

4.1. Insurance Type: {car.insurance_type}
4.2. Insurance Included: {'Yes' if car.insurance_included else 'No'}
4.3. Excess Amount: {car.excess_amount:,.2f} {car.currency}

5. TERMS OF USE

5.1. Mileage Limit: {f'{car.mileage_limit:,} km' if car.mileage_limit else 'Unlimited'}
5.2. Fuel Policy: {car.fuel_policy}

5.3. THE FOLLOWING IS PROHIBITED:
     - Allowing unauthorized persons to drive (unless additional driver specified)
     - Using for taxi or driving instruction purposes
     - Leaving UAE without written consent
     - Smoking inside the vehicle
     - Transporting animals without special permission
     - Driving under influence of alcohol/drugs

5.4. Additional Driver: {'Yes (+' + str(car.additional_driver_fee) + ' ' + car.currency + ')' if car.additional_driver else 'No'}

6. LESSEE'S LIABILITY

6.1. The Lessee is responsible for:
     - Damage not covered by insurance
     - Traffic fines and violations
     - Loss of documents, keys, license plates
     - Interior damage (burns, stains, tears)

6.2. In case of accident, the Lessee must:
     - Immediately notify the Lessor
     - Call police and obtain accident report
     - Not leave the scene of accident

7. VEHICLE RETURN

7.1. The vehicle must be returned:
     - At the specified time and location
     - In clean condition
     - With the same fuel level
     - With all documents and keys

7.2. Late return: penalty equal to daily rate for each day

{f'''8. SPECIAL CONDITIONS

{chr(10).join('- ' + cond for cond in car.special_conditions)}
''' if car.special_conditions else ''}

9. DETAILS AND SIGNATURES

LESSOR:                                   LESSEE:
{company['name_en']}                      {client.full_name_en or client.full_name}
{company['legal_address']}                Passport: {client.passport_number or ''}
Tel: {company['phone']}                   Tel: {client.phone or ''}
Trade License: {company['trade_license']}
IBAN: {company['iban']}

_______________________                   _______________________
{company['director_en']}
Director

Vehicle Handed Over: _______________      Vehicle Received: _______________
Time: _______________                     Time: _______________

VEHICLE RETURN

Date/Time of Return: _______________
Mileage at Return: _______________ km
Fuel Level: _______________
Condition: _______________
Remarks: _______________

Lessor Signature: _______________         Lessee Signature: _______________
"""
        return text

    @staticmethod
    def get_yacht_charter_contract_text(contract: ContractData, language: str = "ru") -> str:
        """Получить текст договора на яхту."""
        if language == "ru":
            return ContractTemplates._yacht_contract_ru(contract)
        else:
            return ContractTemplates._yacht_contract_en(contract)

    @staticmethod
    def _yacht_contract_ru(contract: ContractData) -> str:
        """Текст договора на яхту (русский)."""
        yacht = contract.yacht_data
        client = contract.client
        company = contract.company

        text = f"""
ДОГОВОР ЧАРТЕРА ЯХТЫ
№ {contract.contract_number}

г. Дубай                                                     {contract.created_at}

{company['name_ru']}, именуемое в дальнейшем «Чартерная компания», в лице директора
{company['director']}, с одной стороны, и

{client.full_name}, паспорт {client.passport_number or '________________'},
именуемый(ая) в дальнейшем «Фрахтователь», с другой стороны,

заключили настоящий Договор о нижеследующем:

1. ПРЕДМЕТ ДОГОВОРА

1.1. Чартерная компания предоставляет Фрахтователю яхту:

     Название: {yacht.yacht_name}
     Тип: {yacht.yacht_type}
     Длина: {yacht.yacht_length}
     Вместимость: {yacht.capacity} человек

2. УСЛОВИЯ ЧАРТЕРА

2.1. Дата чартера: {yacht.charter_date}
2.2. Время отправления: {yacht.departure_time}
2.3. Время возвращения: {yacht.return_time}
2.4. Продолжительность: {yacht.duration_hours} часов

2.5. Марина отправления: {yacht.departure_marina}
2.6. Маршрут: {yacht.route_description or 'На усмотрение капитана с учётом погодных условий'}
{f'2.7. Остановки: {", ".join(yacht.stops)}' if yacht.stops else ''}

3. УЧАСТНИКИ

3.1. Количество гостей: {yacht.guests}
3.2. Экипаж включён: {'Да' if yacht.crew_included else 'Нет'}
{f'3.3. Капитан: {yacht.captain_name}' if yacht.captain_name else ''}

4. СТОИМОСТЬ И ОПЛАТА

4.1. Стоимость чартера:
     - Базовая стоимость: {yacht.base_price:,.2f} {yacht.currency}
     - Стоимость за час: {yacht.per_hour_rate:,.2f} {yacht.currency}
     - Общая стоимость: {yacht.total_price:,.2f} {yacht.currency}

4.2. Залоговый депозит: {yacht.deposit_amount:,.2f} {yacht.currency}

4.3. Порядок оплаты:
     - Предоплата {contract.prepayment_percent}%: {contract.prepayment_amount:,.2f} {contract.currency}
     - Остаток не позднее чем за 24 часа до чартера

5. КЕЙТЕРИНГ И ДОПОЛНИТЕЛЬНЫЕ УСЛУГИ

5.1. Кейтеринг включён: {'Да' if yacht.catering_included else 'Нет'}
{f'''5.2. Описание кейтеринга: {yacht.catering_description}
5.3. Стоимость кейтеринга: {yacht.catering_price:,.2f} {yacht.currency}''' if yacht.catering_included else ''}

5.4. Топливо включено: {'Да' if yacht.fuel_included else 'Нет'}
{f'5.5. Топливная политика: {yacht.fuel_policy}' if yacht.fuel_policy else ''}

{f'''6. ОСОБЫЕ ПОЖЕЛАНИЯ

{chr(10).join('- ' + req for req in yacht.special_requests)}
''' if yacht.special_requests else ''}

7. ПРАВИЛА И ОГРАНИЧЕНИЯ

7.1. Фрахтователь и гости обязуются:
     - Соблюдать указания капитана и экипажа
     - Не приносить на борт запрещённые вещества
     - Бережно относиться к имуществу яхты
     - Соблюдать правила безопасности на воде

7.2. ЗАПРЕЩАЕТСЯ:
     - Нахождение на борту лиц, не указанных в договоре
     - Курение в закрытых помещениях
     - Использование открытого огня
     - Прыжки в воду в небезопасных местах

8. ОТМЕНА И ФОРС-МАЖОР

8.1. Условия отмены:
{ContractTemplates._format_cancellation_policy(contract.cancellation_terms or ContractTemplates.CANCELLATION_POLICIES['standard'])}

8.2. В случае неблагоприятных погодных условий:
     - Чартер может быть перенесён на другую дату
     - Полный возврат при невозможности переноса

9. ОТВЕТСТВЕННОСТЬ

9.1. Чартерная компания не несёт ответственности за:
     - Утерю или повреждение личных вещей гостей
     - Травмы, полученные по вине гостей
     - Морскую болезнь

9.2. Фрахтователь несёт ответственность за:
     - Ущерб, причинённый яхте гостями
     - Превышение времени чартера (оплата сверхурочных)

10. РЕКВИЗИТЫ И ПОДПИСИ

ЧАРТЕРНАЯ КОМПАНИЯ:                       ФРАХТОВАТЕЛЬ:
{company['name_ru']}                      {client.full_name}
{company['legal_address']}                {client.address or ''}
Тел: {company['phone']}                   Тел: {client.phone or ''}
Email: {company['email']}                 Email: {client.email or ''}
Trade License: {company['trade_license']}
IBAN: {company['iban']}

_______________________                   _______________________
{company['director']}
Директор

Дата: _______________                     Дата: _______________
"""
        return text

    @staticmethod
    def _yacht_contract_en(contract: ContractData) -> str:
        """Текст договора на яхту (английский)."""
        yacht = contract.yacht_data
        client = contract.client
        company = contract.company

        text = f"""
YACHT CHARTER AGREEMENT
No. {contract.contract_number}

Dubai                                                        {contract.created_at}

{company['name_en']}, hereinafter referred to as the "Charter Company",
represented by Director {company['director_en']}, on the one hand, and

{client.full_name_en or client.full_name}, passport {client.passport_number or '________________'},
hereinafter referred to as the "Charterer", on the other hand,

have concluded this Agreement as follows:

1. SUBJECT OF THE AGREEMENT

1.1. The Charter Company provides the Charterer with the following yacht:

     Name: {yacht.yacht_name}
     Type: {yacht.yacht_type}
     Length: {yacht.yacht_length}
     Capacity: {yacht.capacity} persons

2. CHARTER TERMS

2.1. Charter Date: {yacht.charter_date}
2.2. Departure Time: {yacht.departure_time}
2.3. Return Time: {yacht.return_time}
2.4. Duration: {yacht.duration_hours} hours

2.5. Departure Marina: {yacht.departure_marina}
2.6. Route: {yacht.route_description or 'At captain discretion based on weather conditions'}
{f'2.7. Stops: {", ".join(yacht.stops)}' if yacht.stops else ''}

3. PARTICIPANTS

3.1. Number of Guests: {yacht.guests}
3.2. Crew Included: {'Yes' if yacht.crew_included else 'No'}
{f'3.3. Captain: {yacht.captain_name}' if yacht.captain_name else ''}

4. PRICE AND PAYMENT

4.1. Charter Price:
     - Base Price: {yacht.base_price:,.2f} {yacht.currency}
     - Hourly Rate: {yacht.per_hour_rate:,.2f} {yacht.currency}
     - Total Price: {yacht.total_price:,.2f} {yacht.currency}

4.2. Security Deposit: {yacht.deposit_amount:,.2f} {yacht.currency}

4.3. Payment Terms:
     - Advance Payment {contract.prepayment_percent}%: {contract.prepayment_amount:,.2f} {contract.currency}
     - Balance no later than 24 hours before charter

5. CATERING AND ADDITIONAL SERVICES

5.1. Catering Included: {'Yes' if yacht.catering_included else 'No'}
{f'''5.2. Catering Description: {yacht.catering_description}
5.3. Catering Price: {yacht.catering_price:,.2f} {yacht.currency}''' if yacht.catering_included else ''}

5.4. Fuel Included: {'Yes' if yacht.fuel_included else 'No'}
{f'5.5. Fuel Policy: {yacht.fuel_policy}' if yacht.fuel_policy else ''}

{f'''6. SPECIAL REQUESTS

{chr(10).join('- ' + req for req in yacht.special_requests)}
''' if yacht.special_requests else ''}

7. RULES AND RESTRICTIONS

7.1. The Charterer and guests agree to:
     - Follow captain's and crew's instructions
     - Not bring prohibited substances on board
     - Handle yacht property with care
     - Observe water safety rules

7.2. THE FOLLOWING IS PROHIBITED:
     - Persons not listed in the agreement on board
     - Smoking in enclosed areas
     - Use of open flames
     - Jumping into water in unsafe areas

8. CANCELLATION AND FORCE MAJEURE

8.1. Cancellation Terms:
{ContractTemplates._format_cancellation_policy_en(contract.cancellation_terms or ContractTemplates.CANCELLATION_POLICIES['standard'])}

8.2. In case of adverse weather conditions:
     - Charter may be rescheduled
     - Full refund if rescheduling is not possible

9. LIABILITY

9.1. The Charter Company is not liable for:
     - Loss or damage to guests' personal belongings
     - Injuries caused by guests' own actions
     - Seasickness

9.2. The Charterer is liable for:
     - Damage to the yacht caused by guests
     - Overtime charges for exceeding charter duration

10. DETAILS AND SIGNATURES

CHARTER COMPANY:                          CHARTERER:
{company['name_en']}                      {client.full_name_en or client.full_name}
{company['legal_address']}                {client.address or ''}
Tel: {company['phone']}                   Tel: {client.phone or ''}
Email: {company['email']}                 Email: {client.email or ''}
Trade License: {company['trade_license']}
IBAN: {company['iban']}

_______________________                   _______________________
{company['director_en']}
Director

Date: _______________                     Date: _______________
"""
        return text

    @staticmethod
    def get_agent_contract_text(contract: ContractData, language: str = "ru") -> str:
        """Получить текст агентского договора."""
        if language == "ru":
            return ContractTemplates._agent_contract_ru(contract)
        else:
            return ContractTemplates._agent_contract_en(contract)

    @staticmethod
    def _agent_contract_ru(contract: ContractData) -> str:
        """Текст агентского договора (русский)."""
        agent = contract.agent_data
        company = contract.company

        text = f"""
АГЕНТСКИЙ ДОГОВОР
№ {contract.contract_number}

г. Дубай                                                     {contract.created_at}

{company['name_ru']}, именуемое в дальнейшем «Принципал», в лице директора
{company['director']}, действующего на основании Устава, с одной стороны, и

{agent.agency_name}, Trade License: {agent.trade_license}, в лице
{agent.contact_person}, {agent.contact_position},
именуемое в дальнейшем «Агент», с другой стороны,

заключили настоящий Договор о нижеследующем:

1. ПРЕДМЕТ ДОГОВОРА

1.1. Принципал поручает, а Агент принимает на себя обязательство за вознаграждение
     осуществлять продажу туристических услуг Принципала от своего имени,
     но за счёт Принципала.

1.2. Услуги, охватываемые настоящим договором:
{chr(10).join('     - ' + service for service in agent.services_covered) if agent.services_covered else '     - Все туристические услуги Принципала'}

1.3. Эксклюзивность: {'Да' if agent.exclusive else 'Нет'}
{f'1.4. Территория: {agent.territory}' if agent.territory else ''}

2. КОМИССИОННЫЕ УСЛОВИЯ

2.1. Тип комиссии: {agent.commission_type}
2.2. Размер комиссии: {agent.commission_rate}%

2.3. Комиссия начисляется:
     - На общую стоимость забронированных услуг
     - После полной оплаты клиентом
     - В течение {agent.payment_period} дней после оказания услуги

3. ПОРЯДОК РАСЧЁТОВ

3.1. Тип оплаты: {agent.payment_terms}

3.2. Агент перечисляет Принципалу:
     - 100% стоимости услуг за вычетом комиссии
     - В течение {agent.payment_period} дней с момента получения оплаты от клиента

3.3. Банковские реквизиты Принципала:
     Банк: {company['bank_name']}
     IBAN: {company['iban']}
     SWIFT: {company['swift']}

3.4. Банковские реквизиты Агента:
     Банк: {agent.bank_name}
     IBAN: {agent.iban}
     SWIFT: {agent.swift}

4. ПРАВА И ОБЯЗАННОСТИ ПРИНЦИПАЛА

4.1. Принципал обязуется:
     - Предоставлять актуальную информацию о ценах и услугах
     - Своевременно подтверждать бронирования
     - Качественно оказывать забронированные услуги
     - Предоставлять рекламные материалы
     - Уведомлять об изменениях за 7 дней

4.2. Принципал имеет право:
     - Устанавливать и изменять цены
     - Отказать в подтверждении бронирования
     - Контролировать деятельность Агента
     - Расторгнуть договор при нарушении условий

5. ПРАВА И ОБЯЗАННОСТИ АГЕНТА

5.1. Агент обязуется:
     - Активно продвигать услуги Принципала
     - Своевременно передавать бронирования
     - Предоставлять полную информацию о клиентах
     - Соблюдать ценовую политику Принципала
     - Вести учёт продаж
     - Передавать оплату в установленные сроки

5.2. Агент имеет право:
     - Получать комиссионное вознаграждение
     - Использовать рекламные материалы Принципала
     - Получать приоритетное подтверждение бронирований

6. СРОК ДЕЙСТВИЯ ДОГОВОРА

6.1. Дата начала: {agent.contract_start}
6.2. Дата окончания: {agent.contract_end}
6.3. Автоматическое продление: {'Да' if agent.auto_renewal else 'Нет'}

6.4. Уведомление о расторжении: за {agent.notice_period} дней

7. КОНФИДЕНЦИАЛЬНОСТЬ

7.1. Стороны обязуются сохранять конфиденциальность:
     - Коммерческих условий договора
     - Клиентской базы
     - Внутренних цен и ставок

8. ОТВЕТСТВЕННОСТЬ

8.1. За неисполнение обязательств стороны несут ответственность
     в соответствии с законодательством ОАЭ.

8.2. Принципал не несёт ответственности за действия Агента
     перед третьими лицами.

9. РАЗРЕШЕНИЕ СПОРОВ

9.1. Споры решаются путём переговоров.
9.2. При недостижении согласия - в суде ОАЭ.

10. РЕКВИЗИТЫ И ПОДПИСИ

ПРИНЦИПАЛ:                                АГЕНТ:
{company['name_ru']}                      {agent.agency_name}
{company['legal_address']}                {agent.legal_address}
Trade License: {company['trade_license']} Trade License: {agent.trade_license}
Тел: {company['phone']}                   Тел: {agent.contact_phone}
Email: {company['email']}                 Email: {agent.contact_email}
IBAN: {company['iban']}                   IBAN: {agent.iban}

_______________________                   _______________________
{company['director']}                     {agent.contact_person}
Директор                                  {agent.contact_position}

Дата: _______________                     Дата: _______________
М.П.                                      М.П.
"""
        return text

    @staticmethod
    def _agent_contract_en(contract: ContractData) -> str:
        """Текст агентского договора (английский)."""
        agent = contract.agent_data
        company = contract.company

        text = f"""
AGENCY AGREEMENT
No. {contract.contract_number}

Dubai                                                        {contract.created_at}

{company['name_en']}, hereinafter referred to as the "Principal",
represented by Director {company['director_en']}, on the one hand, and

{agent.agency_name_en or agent.agency_name}, Trade License: {agent.trade_license},
represented by {agent.contact_person}, {agent.contact_position},
hereinafter referred to as the "Agent", on the other hand,

have concluded this Agreement as follows:

1. SUBJECT OF THE AGREEMENT

1.1. The Principal engages the Agent, and the Agent accepts the engagement,
     to sell Principal's tourism services on Agent's behalf but at Principal's expense.

1.2. Services covered by this Agreement:
{chr(10).join('     - ' + service for service in agent.services_covered) if agent.services_covered else '     - All Principal tourism services'}

1.3. Exclusivity: {'Yes' if agent.exclusive else 'No'}
{f'1.4. Territory: {agent.territory}' if agent.territory else ''}

2. COMMISSION TERMS

2.1. Commission Type: {agent.commission_type}
2.2. Commission Rate: {agent.commission_rate}%

2.3. Commission is calculated:
     - On the total value of booked services
     - After full payment by the client
     - Within {agent.payment_period} days after service delivery

3. PAYMENT PROCEDURES

3.1. Payment Type: {agent.payment_terms}

3.2. The Agent shall transfer to the Principal:
     - 100% of service cost less commission
     - Within {agent.payment_period} days of receiving client payment

3.3. Principal's Banking Details:
     Bank: {company['bank_name']}
     IBAN: {company['iban']}
     SWIFT: {company['swift']}

3.4. Agent's Banking Details:
     Bank: {agent.bank_name}
     IBAN: {agent.iban}
     SWIFT: {agent.swift}

4. PRINCIPAL'S RIGHTS AND OBLIGATIONS

4.1. The Principal undertakes to:
     - Provide current pricing and service information
     - Confirm bookings promptly
     - Deliver booked services with quality
     - Provide promotional materials
     - Notify of changes 7 days in advance

4.2. The Principal has the right to:
     - Set and modify prices
     - Decline booking confirmations
     - Monitor Agent's activities
     - Terminate agreement upon breach

5. AGENT'S RIGHTS AND OBLIGATIONS

5.1. The Agent undertakes to:
     - Actively promote Principal's services
     - Submit bookings promptly
     - Provide complete client information
     - Adhere to Principal's pricing policy
     - Maintain sales records
     - Transfer payments within specified terms

5.2. The Agent has the right to:
     - Receive commission payments
     - Use Principal's promotional materials
     - Receive priority booking confirmations

6. TERM OF AGREEMENT

6.1. Start Date: {agent.contract_start}
6.2. End Date: {agent.contract_end}
6.3. Auto-renewal: {'Yes' if agent.auto_renewal else 'No'}

6.4. Termination Notice: {agent.notice_period} days

7. CONFIDENTIALITY

7.1. Both parties agree to maintain confidentiality of:
     - Commercial terms of this agreement
     - Client databases
     - Internal pricing and rates

8. LIABILITY

8.1. Both parties are liable for non-performance in accordance with UAE law.

8.2. The Principal is not liable for Agent's actions towards third parties.

9. DISPUTE RESOLUTION

9.1. Disputes shall be resolved through negotiation.
9.2. If no agreement is reached - UAE courts.

10. DETAILS AND SIGNATURES

PRINCIPAL:                                AGENT:
{company['name_en']}                      {agent.agency_name_en or agent.agency_name}
{company['legal_address']}                {agent.legal_address}
Trade License: {company['trade_license']} Trade License: {agent.trade_license}
Tel: {company['phone']}                   Tel: {agent.contact_phone}
Email: {company['email']}                 Email: {agent.contact_email}
IBAN: {company['iban']}                   IBAN: {agent.iban}

_______________________                   _______________________
{company['director_en']}                  {agent.contact_person}
Director                                  {agent.contact_position}

Date: _______________                     Date: _______________
"""
        return text

    @staticmethod
    def _format_cancellation_policy(terms: List[Dict]) -> str:
        """Форматировать условия отмены (русский)."""
        lines = []
        sorted_terms = sorted(terms, key=lambda x: x.get("days_before", 0), reverse=True)

        for term in sorted_terms:
            days = term.get("days_before", 0)
            refund = term.get("refund_percent", 0)

            if days == 0:
                lines.append(f"     - В день оказания услуги: возврат {refund}%")
            elif refund == 100:
                lines.append(f"     - Более {days} дней до услуги: полный возврат")
            elif refund == 0:
                lines.append(f"     - Менее {days} дней до услуги: без возврата")
            else:
                lines.append(f"     - От {days} дней до услуги: возврат {refund}%")

        return "\n".join(lines)

    @staticmethod
    def _format_cancellation_policy_en(terms: List[Dict]) -> str:
        """Форматировать условия отмены (английский)."""
        lines = []
        sorted_terms = sorted(terms, key=lambda x: x.get("days_before", 0), reverse=True)

        for term in sorted_terms:
            days = term.get("days_before", 0)
            refund = term.get("refund_percent", 0)

            if days == 0:
                lines.append(f"     - On the day of service: {refund}% refund")
            elif refund == 100:
                lines.append(f"     - More than {days} days before: full refund")
            elif refund == 0:
                lines.append(f"     - Less than {days} days before: no refund")
            else:
                lines.append(f"     - {days} or more days before: {refund}% refund")

        return "\n".join(lines)


# ═══════════════════════════════════════════════════════════════
# ГЕНЕРАТОР PDF
# ═══════════════════════════════════════════════════════════════

class PDFGenerator:
    """Генератор PDF документов."""

    def __init__(self):
        if not REPORTLAB_AVAILABLE:
            raise ImportError("reportlab library not installed. Run: pip install reportlab")

        self._register_fonts()
        self.styles = getSampleStyleSheet()
        self._setup_styles()

    def _register_fonts(self):
        """Регистрация шрифтов для кириллицы."""
        # Пути к шрифтам
        font_paths = [
            Path("C:/Windows/Fonts"),
            Path("/usr/share/fonts/truetype"),
            Path("/System/Library/Fonts"),
            SCRIPTS_DIR / "fonts"
        ]

        font_registered = False

        for font_path in font_paths:
            if font_path.exists():
                # Пробуем разные шрифты
                fonts_to_try = [
                    ("DejaVuSans.ttf", "DejaVuSans"),
                    ("arial.ttf", "Arial"),
                    ("times.ttf", "Times"),
                    ("calibri.ttf", "Calibri"),
                ]

                for font_file, font_name in fonts_to_try:
                    try:
                        full_path = font_path / font_file
                        if full_path.exists():
                            pdfmetrics.registerFont(TTFont(font_name, str(full_path)))
                            self.default_font = font_name
                            font_registered = True
                            break
                    except:
                        continue

            if font_registered:
                break

        if not font_registered:
            self.default_font = "Helvetica"

    def _setup_styles(self):
        """Настройка стилей документа."""
        self.styles.add(ParagraphStyle(
            name='ContractTitle',
            fontName=self.default_font,
            fontSize=14,
            leading=18,
            alignment=1,  # Center
            spaceAfter=20
        ))

        self.styles.add(ParagraphStyle(
            name='ContractBody',
            fontName=self.default_font,
            fontSize=10,
            leading=14,
            alignment=0,  # Left
            spaceAfter=6
        ))

        self.styles.add(ParagraphStyle(
            name='ContractHeader',
            fontName=self.default_font,
            fontSize=12,
            leading=16,
            alignment=0,
            spaceBefore=12,
            spaceAfter=6,
            fontWeight='bold'
        ))

    def generate(
        self,
        contract: ContractData,
        output_path: Path,
        include_qr: bool = True
    ) -> Path:
        """Генерировать PDF документ."""
        # Создание директории
        output_path.mkdir(parents=True, exist_ok=True)

        # Путь к файлу
        filename = f"{contract.contract_number}_{contract.contract_id}.pdf"
        pdf_path = output_path / filename

        # Создание документа
        doc = SimpleDocTemplate(
            str(pdf_path),
            pagesize=A4,
            rightMargin=2*cm,
            leftMargin=2*cm,
            topMargin=2*cm,
            bottomMargin=2*cm
        )

        # Получение текста договора
        text = self._get_contract_text(contract)

        # Элементы документа
        elements = []

        # Добавление текста
        for line in text.split('\n'):
            if line.strip():
                if line.strip().startswith(('1.', '2.', '3.', '4.', '5.', '6.', '7.', '8.', '9.', '10.')):
                    elements.append(Paragraph(line.strip(), self.styles['ContractHeader']))
                elif 'ДОГОВОР' in line or 'AGREEMENT' in line:
                    elements.append(Paragraph(line.strip(), self.styles['ContractTitle']))
                else:
                    # Замена пробелов на неразрывные для сохранения форматирования
                    formatted_line = line.replace('  ', '&nbsp;&nbsp;')
                    elements.append(Paragraph(formatted_line, self.styles['ContractBody']))
            else:
                elements.append(Spacer(1, 6))

        # QR код
        if include_qr:
            qr_path = QRCodeGenerator.generate(
                contract.contract_id,
                contract.contract_number,
                output_path
            )

            elements.append(Spacer(1, 20))
            elements.append(Paragraph("Для проверки подлинности:", self.styles['ContractBody']))
            elements.append(Image(str(qr_path), width=3*cm, height=3*cm))

        # Сборка документа
        doc.build(elements)

        return pdf_path

    def _get_contract_text(self, contract: ContractData) -> str:
        """Получить текст договора в зависимости от типа."""
        if contract.contract_type == ContractType.TOUR:
            return ContractTemplates.get_tour_contract_text(contract, contract.language)
        elif contract.contract_type == ContractType.CAR_RENTAL:
            return ContractTemplates.get_car_rental_contract_text(contract, contract.language)
        elif contract.contract_type == ContractType.YACHT:
            return ContractTemplates.get_yacht_charter_contract_text(contract, contract.language)
        elif contract.contract_type == ContractType.AGENT:
            return ContractTemplates.get_agent_contract_text(contract, contract.language)
        else:
            return "Тип договора не поддерживается"


# ═══════════════════════════════════════════════════════════════
# ГЕНЕРАТОР DOCX
# ═══════════════════════════════════════════════════════════════

class DOCXGenerator:
    """Генератор DOCX документов."""

    def __init__(self):
        if not DOCX_AVAILABLE:
            raise ImportError("python-docx library not installed. Run: pip install python-docx")

    def generate(
        self,
        contract: ContractData,
        output_path: Path,
        include_qr: bool = True
    ) -> Path:
        """Генерировать DOCX документ."""
        # Создание директории
        output_path.mkdir(parents=True, exist_ok=True)

        # Путь к файлу
        filename = f"{contract.contract_number}_{contract.contract_id}.docx"
        docx_path = output_path / filename

        # Создание документа
        doc = Document()

        # Настройка полей
        sections = doc.sections
        for section in sections:
            section.top_margin = Cm(2)
            section.bottom_margin = Cm(2)
            section.left_margin = Cm(2)
            section.right_margin = Cm(2)

        # Получение текста договора
        text = self._get_contract_text(contract)

        # Добавление текста
        for line in text.split('\n'):
            if line.strip():
                para = doc.add_paragraph()

                if 'ДОГОВОР' in line or 'AGREEMENT' in line:
                    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    run = para.add_run(line.strip())
                    run.bold = True
                    run.font.size = Pt(14)
                elif line.strip().startswith(('1.', '2.', '3.', '4.', '5.', '6.', '7.', '8.', '9.', '10.')):
                    run = para.add_run(line.strip())
                    run.bold = True
                    run.font.size = Pt(11)
                else:
                    run = para.add_run(line)
                    run.font.size = Pt(10)

        # QR код
        if include_qr:
            qr_path = QRCodeGenerator.generate(
                contract.contract_id,
                contract.contract_number,
                output_path
            )

            doc.add_paragraph()
            para = doc.add_paragraph("Для проверки подлинности:")
            doc.add_picture(str(qr_path), width=Cm(3))

        # Сохранение
        doc.save(str(docx_path))

        return docx_path

    def _get_contract_text(self, contract: ContractData) -> str:
        """Получить текст договора в зависимости от типа."""
        if contract.contract_type == ContractType.TOUR:
            return ContractTemplates.get_tour_contract_text(contract, contract.language)
        elif contract.contract_type == ContractType.CAR_RENTAL:
            return ContractTemplates.get_car_rental_contract_text(contract, contract.language)
        elif contract.contract_type == ContractType.YACHT:
            return ContractTemplates.get_yacht_charter_contract_text(contract, contract.language)
        elif contract.contract_type == ContractType.AGENT:
            return ContractTemplates.get_agent_contract_text(contract, contract.language)
        else:
            return "Тип договора не поддерживается"


# ═══════════════════════════════════════════════════════════════
# ГЛАВНЫЙ КЛАСС ГЕНЕРАТОРА ДОГОВОРОВ
# ═══════════════════════════════════════════════════════════════

class ContractGenerator:
    """
    Главный класс для генерации договоров.

    Пример использования:

    ```python
    generator = ContractGenerator()

    # Создание данных клиента
    client = ClientData(
        full_name="Иванов Иван Иванович",
        passport_number="123456789",
        phone="+7 999 123 45 67",
        email="ivanov@example.com"
    )

    # Создание данных тура
    tour = TourServiceData(
        name="Обзорная экскурсия по Дубаю",
        date="15.02.2025",
        time="09:00",
        duration="8 часов",
        participants=2,
        price_per_person=150,
        total_price=300,
        currency="AED",
        pickup_location="Ваш отель"
    )

    # Генерация договора
    contract = generator.create_tour_contract(
        client=client,
        tour_data=tour,
        prepayment_percent=50
    )

    # Генерация документов
    pdf_path = generator.generate_pdf(contract)
    docx_path = generator.generate_docx(contract)
    ```
    """

    def __init__(
        self,
        contracts_dir: Path = CONTRACTS_DIR,
        company: Dict[str, str] = None
    ):
        self.contracts_dir = contracts_dir
        self.company = company or DEFAULT_COMPANY.copy()

        # Инициализация компонентов
        self.number_generator = ContractNumberGenerator()
        self.registry = ContractRegistry()

        # Генераторы документов (ленивая инициализация)
        self._pdf_generator = None
        self._docx_generator = None
        self._signer = None

        # Создание директорий
        self._ensure_directories()

    def _ensure_directories(self):
        """Создание необходимых директорий."""
        self.contracts_dir.mkdir(parents=True, exist_ok=True)
        (self.contracts_dir / "pdf").mkdir(exist_ok=True)
        (self.contracts_dir / "docx").mkdir(exist_ok=True)
        (self.contracts_dir / "qr").mkdir(exist_ok=True)

    @property
    def pdf_generator(self) -> PDFGenerator:
        """Ленивая инициализация PDF генератора."""
        if self._pdf_generator is None:
            self._pdf_generator = PDFGenerator()
        return self._pdf_generator

    @property
    def docx_generator(self) -> DOCXGenerator:
        """Ленивая инициализация DOCX генератора."""
        if self._docx_generator is None:
            self._docx_generator = DOCXGenerator()
        return self._docx_generator

    # ═══════════════════════════════════════════════════════════════
    # СОЗДАНИЕ ДОГОВОРОВ
    # ═══════════════════════════════════════════════════════════════

    def create_tour_contract(
        self,
        client: ClientData,
        tour_data: TourServiceData,
        prepayment_percent: float = 50,
        cancellation_policy: str = "standard",
        language: str = "ru",
        notes: str = None
    ) -> ContractData:
        """Создать договор на тур/экскурсию."""
        contract_id = self.number_generator.generate_id()
        contract_number = self.number_generator.generate(ContractType.TOUR)

        now = datetime.now()

        contract = ContractData(
            contract_id=contract_id,
            contract_type=ContractType.TOUR,
            contract_number=contract_number,
            created_at=now.strftime("%d.%m.%Y"),
            valid_from=now.strftime("%d.%m.%Y"),
            valid_until=tour_data.date,
            status=ContractStatus.DRAFT,
            client=client,
            company=self.company,
            tour_data=tour_data,
            total_amount=tour_data.total_price,
            currency=tour_data.currency,
            prepayment_percent=prepayment_percent,
            prepayment_amount=tour_data.total_price * prepayment_percent / 100,
            cancellation_policy=cancellation_policy,
            cancellation_terms=ContractTemplates.CANCELLATION_POLICIES.get(cancellation_policy, []),
            language=language,
            notes=notes
        )

        # Добавление в реестр
        self.registry.add_contract(contract)

        return contract

    def create_car_rental_contract(
        self,
        client: ClientData,
        car_data: CarRentalData,
        prepayment_percent: float = 100,
        language: str = "ru",
        notes: str = None
    ) -> ContractData:
        """Создать договор аренды авто."""
        contract_id = self.number_generator.generate_id()
        contract_number = self.number_generator.generate(ContractType.CAR_RENTAL)

        now = datetime.now()

        contract = ContractData(
            contract_id=contract_id,
            contract_type=ContractType.CAR_RENTAL,
            contract_number=contract_number,
            created_at=now.strftime("%d.%m.%Y"),
            valid_from=car_data.rental_start,
            valid_until=car_data.rental_end,
            status=ContractStatus.DRAFT,
            client=client,
            company=self.company,
            car_rental_data=car_data,
            total_amount=car_data.total_price,
            currency=car_data.currency,
            prepayment_percent=prepayment_percent,
            prepayment_amount=car_data.total_price * prepayment_percent / 100,
            language=language,
            notes=notes
        )

        self.registry.add_contract(contract)

        return contract

    def create_yacht_charter_contract(
        self,
        client: ClientData,
        yacht_data: YachtCharterData,
        prepayment_percent: float = 50,
        cancellation_policy: str = "standard",
        language: str = "ru",
        notes: str = None
    ) -> ContractData:
        """Создать договор чартера яхты."""
        contract_id = self.number_generator.generate_id()
        contract_number = self.number_generator.generate(ContractType.YACHT)

        now = datetime.now()

        contract = ContractData(
            contract_id=contract_id,
            contract_type=ContractType.YACHT,
            contract_number=contract_number,
            created_at=now.strftime("%d.%m.%Y"),
            valid_from=now.strftime("%d.%m.%Y"),
            valid_until=yacht_data.charter_date,
            status=ContractStatus.DRAFT,
            client=client,
            company=self.company,
            yacht_data=yacht_data,
            total_amount=yacht_data.total_price,
            currency=yacht_data.currency,
            prepayment_percent=prepayment_percent,
            prepayment_amount=yacht_data.total_price * prepayment_percent / 100,
            cancellation_policy=cancellation_policy,
            cancellation_terms=ContractTemplates.CANCELLATION_POLICIES.get(cancellation_policy, []),
            language=language,
            notes=notes
        )

        self.registry.add_contract(contract)

        return contract

    def create_agent_contract(
        self,
        agent_data: AgentContractData,
        language: str = "ru",
        notes: str = None
    ) -> ContractData:
        """Создать агентский договор."""
        contract_id = self.number_generator.generate_id()
        contract_number = self.number_generator.generate(ContractType.AGENT)

        now = datetime.now()

        # Создаем ClientData из данных агентства
        client = ClientData(
            full_name=agent_data.contact_person,
            company_name=agent_data.agency_name,
            company_name_en=agent_data.agency_name_en,
            trade_license=agent_data.trade_license,
            vat_number=agent_data.vat_number,
            legal_address=agent_data.legal_address,
            phone=agent_data.contact_phone,
            email=agent_data.contact_email,
            position=agent_data.contact_position,
            bank_name=agent_data.bank_name,
            iban=agent_data.iban,
            swift=agent_data.swift
        )

        contract = ContractData(
            contract_id=contract_id,
            contract_type=ContractType.AGENT,
            contract_number=contract_number,
            created_at=now.strftime("%d.%m.%Y"),
            valid_from=agent_data.contract_start,
            valid_until=agent_data.contract_end,
            status=ContractStatus.DRAFT,
            client=client,
            company=self.company,
            agent_data=agent_data,
            total_amount=0,  # Агентский договор без фиксированной суммы
            currency="AED",
            language=language,
            notes=notes
        )

        self.registry.add_contract(contract)

        return contract

    # ═══════════════════════════════════════════════════════════════
    # ГЕНЕРАЦИЯ ДОКУМЕНТОВ
    # ═══════════════════════════════════════════════════════════════

    def generate_pdf(
        self,
        contract: ContractData,
        include_qr: bool = True
    ) -> Path:
        """Генерировать PDF версию договора."""
        output_dir = self.contracts_dir / "pdf"
        return self.pdf_generator.generate(contract, output_dir, include_qr)

    def generate_docx(
        self,
        contract: ContractData,
        include_qr: bool = True
    ) -> Path:
        """Генерировать DOCX версию договора."""
        output_dir = self.contracts_dir / "docx"
        return self.docx_generator.generate(contract, output_dir, include_qr)

    def generate_both(
        self,
        contract: ContractData,
        include_qr: bool = True
    ) -> tuple:
        """Генерировать PDF и DOCX версии."""
        pdf_path = self.generate_pdf(contract, include_qr)
        docx_path = self.generate_docx(contract, include_qr)
        return pdf_path, docx_path

    def generate_text(self, contract: ContractData) -> str:
        """Получить текстовую версию договора."""
        if contract.contract_type == ContractType.TOUR:
            return ContractTemplates.get_tour_contract_text(contract, contract.language)
        elif contract.contract_type == ContractType.CAR_RENTAL:
            return ContractTemplates.get_car_rental_contract_text(contract, contract.language)
        elif contract.contract_type == ContractType.YACHT:
            return ContractTemplates.get_yacht_charter_contract_text(contract, contract.language)
        elif contract.contract_type == ContractType.AGENT:
            return ContractTemplates.get_agent_contract_text(contract, contract.language)
        return ""

    # ═══════════════════════════════════════════════════════════════
    # УПРАВЛЕНИЕ ДОГОВОРАМИ
    # ═══════════════════════════════════════════════════════════════

    def update_status(
        self,
        contract: ContractData,
        new_status: ContractStatus
    ) -> bool:
        """Обновить статус договора."""
        contract.status = new_status

        if new_status == ContractStatus.SIGNED:
            contract.signed_at = datetime.now().strftime("%d.%m.%Y %H:%M")

        return self.registry.update_contract(contract)

    def get_contract(self, contract_id: str) -> Optional[Dict]:
        """Получить договор по ID."""
        return self.registry.get_contract(contract_id)

    def find_contracts(self, **kwargs) -> List[Dict]:
        """Поиск договоров по критериям."""
        return self.registry.find_contracts(**kwargs)

    def get_statistics(self) -> Dict:
        """Получить статистику по договорам."""
        return self.registry.get_statistics()

    # ═══════════════════════════════════════════════════════════════
    # АВТОЗАПОЛНЕНИЕ ИЗ ДАННЫХ КЛИЕНТА
    # ═══════════════════════════════════════════════════════════════

    @staticmethod
    def client_from_dict(data: Dict) -> ClientData:
        """Создать ClientData из словаря."""
        return ClientData(
            full_name=data.get("full_name", data.get("name", "")),
            full_name_en=data.get("full_name_en", data.get("name_en")),
            passport_number=data.get("passport_number", data.get("passport")),
            passport_issued=data.get("passport_issued"),
            passport_expiry=data.get("passport_expiry"),
            nationality=data.get("nationality"),
            date_of_birth=data.get("date_of_birth", data.get("dob")),
            phone=data.get("phone", data.get("tel")),
            email=data.get("email"),
            address=data.get("address"),
            company_name=data.get("company_name", data.get("company")),
            company_name_en=data.get("company_name_en"),
            trade_license=data.get("trade_license"),
            vat_number=data.get("vat_number", data.get("vat")),
            legal_address=data.get("legal_address"),
            bank_name=data.get("bank_name", data.get("bank")),
            iban=data.get("iban"),
            swift=data.get("swift"),
            authorized_person=data.get("authorized_person"),
            position=data.get("position")
        )

    @staticmethod
    def tour_from_dict(data: Dict) -> TourServiceData:
        """Создать TourServiceData из словаря."""
        return TourServiceData(
            name=data.get("name", data.get("tour_name", "")),
            name_en=data.get("name_en"),
            description=data.get("description"),
            date=data.get("date", ""),
            time=data.get("time", ""),
            duration=data.get("duration", ""),
            pickup_location=data.get("pickup_location", data.get("pickup")),
            dropoff_location=data.get("dropoff_location", data.get("dropoff")),
            participants=int(data.get("participants", data.get("adults", 1))),
            children=int(data.get("children", 0)),
            price_per_person=float(data.get("price_per_person", data.get("pax_price", 0))),
            total_price=float(data.get("total_price", data.get("total", 0))),
            currency=data.get("currency", "AED"),
            includes=data.get("includes", []),
            excludes=data.get("excludes", []),
            notes=data.get("notes")
        )

    @staticmethod
    def car_rental_from_dict(data: Dict) -> CarRentalData:
        """Создать CarRentalData из словаря."""
        return CarRentalData(
            vehicle_make=data.get("make", data.get("vehicle_make", "")),
            vehicle_model=data.get("model", data.get("vehicle_model", "")),
            vehicle_year=int(data.get("year", data.get("vehicle_year", 2024))),
            plate_number=data.get("plate", data.get("plate_number", "")),
            vin_number=data.get("vin", data.get("vin_number")),
            color=data.get("color"),
            mileage_start=int(data.get("mileage", data.get("mileage_start", 0))),
            mileage_limit=data.get("mileage_limit"),
            rental_start=data.get("start_date", data.get("rental_start", "")),
            rental_end=data.get("end_date", data.get("rental_end", "")),
            pickup_location=data.get("pickup", data.get("pickup_location", "")),
            dropoff_location=data.get("dropoff", data.get("dropoff_location", "")),
            daily_rate=float(data.get("daily_rate", data.get("rate", 0))),
            weekly_rate=data.get("weekly_rate"),
            monthly_rate=data.get("monthly_rate"),
            total_days=int(data.get("days", data.get("total_days", 1))),
            total_price=float(data.get("total", data.get("total_price", 0))),
            currency=data.get("currency", "AED"),
            deposit_amount=float(data.get("deposit", data.get("deposit_amount", 0))),
            deposit_type=data.get("deposit_type", "cash"),
            insurance_type=data.get("insurance_type", "basic"),
            insurance_included=data.get("insurance_included", True),
            excess_amount=float(data.get("excess", data.get("excess_amount", 0))),
            additional_driver=data.get("additional_driver", False),
            additional_driver_fee=float(data.get("additional_driver_fee", 0)),
            fuel_policy=data.get("fuel_policy", "full_to_full"),
            special_conditions=data.get("special_conditions", [])
        )

    @staticmethod
    def yacht_from_dict(data: Dict) -> YachtCharterData:
        """Создать YachtCharterData из словаря."""
        return YachtCharterData(
            yacht_name=data.get("name", data.get("yacht_name", "")),
            yacht_type=data.get("type", data.get("yacht_type", "motor")),
            yacht_length=data.get("length", data.get("yacht_length", "")),
            capacity=int(data.get("capacity", 10)),
            charter_date=data.get("date", data.get("charter_date", "")),
            departure_time=data.get("departure", data.get("departure_time", "")),
            return_time=data.get("return", data.get("return_time", "")),
            duration_hours=int(data.get("hours", data.get("duration_hours", 4))),
            departure_marina=data.get("marina", data.get("departure_marina", "")),
            route_description=data.get("route", data.get("route_description")),
            stops=data.get("stops", []),
            guests=int(data.get("guests", 1)),
            crew_included=data.get("crew_included", True),
            captain_name=data.get("captain", data.get("captain_name")),
            base_price=float(data.get("base_price", 0)),
            per_hour_rate=float(data.get("hourly_rate", data.get("per_hour_rate", 0))),
            total_price=float(data.get("total", data.get("total_price", 0))),
            currency=data.get("currency", "AED"),
            deposit_amount=float(data.get("deposit", data.get("deposit_amount", 0))),
            catering_included=data.get("catering_included", False),
            catering_description=data.get("catering_description"),
            catering_price=float(data.get("catering_price", 0)),
            fuel_included=data.get("fuel_included", True),
            fuel_policy=data.get("fuel_policy"),
            special_requests=data.get("special_requests", [])
        )

    @staticmethod
    def agent_from_dict(data: Dict) -> AgentContractData:
        """Создать AgentContractData из словаря."""
        return AgentContractData(
            agency_name=data.get("name", data.get("agency_name", "")),
            agency_name_en=data.get("name_en", data.get("agency_name_en")),
            trade_license=data.get("license", data.get("trade_license", "")),
            vat_number=data.get("vat", data.get("vat_number")),
            legal_address=data.get("address", data.get("legal_address", "")),
            contact_person=data.get("contact", data.get("contact_person", "")),
            contact_position=data.get("position", data.get("contact_position", "")),
            contact_phone=data.get("phone", data.get("contact_phone", "")),
            contact_email=data.get("email", data.get("contact_email", "")),
            commission_rate=float(data.get("commission", data.get("commission_rate", 10))),
            commission_type=data.get("commission_type", "percentage"),
            payment_terms=data.get("payment_terms", "prepaid"),
            payment_period=int(data.get("payment_period", 7)),
            contract_start=data.get("start", data.get("contract_start", "")),
            contract_end=data.get("end", data.get("contract_end", "")),
            auto_renewal=data.get("auto_renewal", True),
            notice_period=int(data.get("notice_period", 30)),
            services_covered=data.get("services", data.get("services_covered", [])),
            exclusive=data.get("exclusive", False),
            territory=data.get("territory"),
            bank_name=data.get("bank", data.get("bank_name", "")),
            iban=data.get("iban", ""),
            swift=data.get("swift", "")
        )


# ═══════════════════════════════════════════════════════════════
# CLI ИНТЕРФЕЙС
# ═══════════════════════════════════════════════════════════════

def main():
    """Главная функция CLI."""
    import argparse

    parser = argparse.ArgumentParser(
        description="Генератор договоров с автозаполнением",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
Примеры использования:

  # Создать тестовый договор на тур
  python contract_generator.py --demo tour

  # Создать договор из JSON файла
  python contract_generator.py --input client_data.json --type tour

  # Показать статистику
  python contract_generator.py --stats

  # Найти договоры клиента
  python contract_generator.py --find "Иванов"
        """
    )

    parser.add_argument(
        "--demo",
        choices=["tour", "car", "yacht", "agent"],
        help="Создать демо-договор указанного типа"
    )

    parser.add_argument(
        "--input",
        type=str,
        help="JSON файл с данными для договора"
    )

    parser.add_argument(
        "--type",
        choices=["tour", "car", "yacht", "agent"],
        help="Тип договора (требуется с --input)"
    )

    parser.add_argument(
        "--output",
        type=str,
        default=str(CONTRACTS_DIR),
        help=f"Папка для сохранения (по умолчанию: {CONTRACTS_DIR})"
    )

    parser.add_argument(
        "--format",
        choices=["pdf", "docx", "both", "text"],
        default="both",
        help="Формат вывода (по умолчанию: both)"
    )

    parser.add_argument(
        "--lang",
        choices=["ru", "en"],
        default="ru",
        help="Язык договора (по умолчанию: ru)"
    )

    parser.add_argument(
        "--stats",
        action="store_true",
        help="Показать статистику договоров"
    )

    parser.add_argument(
        "--find",
        type=str,
        help="Найти договоры по имени клиента"
    )

    parser.add_argument(
        "--no-qr",
        action="store_true",
        help="Не добавлять QR код"
    )

    args = parser.parse_args()

    generator = ContractGenerator()

    # Статистика
    if args.stats:
        stats = generator.get_statistics()
        print("\n=== Статистика договоров ===")
        print(f"Всего договоров: {stats['total']}")
        print(f"За этот месяц: {stats['this_month']}")
        print(f"За этот год: {stats['this_year']}")
        print(f"Общая сумма: {stats['total_amount']:,.2f} AED")
        print("\nПо типам:")
        for t, count in stats['by_type'].items():
            print(f"  {t}: {count}")
        print("\nПо статусам:")
        for s, count in stats['by_status'].items():
            print(f"  {s}: {count}")
        return

    # Поиск
    if args.find:
        contracts = generator.find_contracts(client_name=args.find)
        print(f"\n=== Найдено договоров: {len(contracts)} ===")
        for c in contracts:
            client = c.get('client', {})
            print(f"  {c['contract_number']} | {client.get('full_name', 'N/A')} | "
                  f"{c['contract_type']} | {c['status']} | {c.get('total_amount', 0):,.2f} {c.get('currency', 'AED')}")
        return

    # Демо
    if args.demo:
        contract = create_demo_contract(generator, args.demo, args.lang)
        output_contract(generator, contract, args.format, not args.no_qr)
        return

    # Из файла
    if args.input and args.type:
        with open(args.input, 'r', encoding='utf-8') as f:
            data = json.load(f)

        client = generator.client_from_dict(data.get('client', {}))

        if args.type == "tour":
            tour_data = generator.tour_from_dict(data.get('tour', data.get('service', {})))
            contract = generator.create_tour_contract(
                client=client,
                tour_data=tour_data,
                language=args.lang
            )
        elif args.type == "car":
            car_data = generator.car_rental_from_dict(data.get('car', data.get('rental', {})))
            contract = generator.create_car_rental_contract(
                client=client,
                car_data=car_data,
                language=args.lang
            )
        elif args.type == "yacht":
            yacht_data = generator.yacht_from_dict(data.get('yacht', data.get('charter', {})))
            contract = generator.create_yacht_charter_contract(
                client=client,
                yacht_data=yacht_data,
                language=args.lang
            )
        elif args.type == "agent":
            agent_data = generator.agent_from_dict(data.get('agent', data.get('agency', {})))
            contract = generator.create_agent_contract(
                agent_data=agent_data,
                language=args.lang
            )
        else:
            print("Неизвестный тип договора")
            return

        output_contract(generator, contract, args.format, not args.no_qr)
        return

    parser.print_help()


def create_demo_contract(generator: ContractGenerator, demo_type: str, language: str) -> ContractData:
    """Создать демо-договор."""

    if demo_type == "tour":
        client = ClientData(
            full_name="Иванов Иван Иванович",
            full_name_en="Ivan Ivanov",
            passport_number="75 1234567",
            phone="+7 999 123 45 67",
            email="ivanov@example.com"
        )

        tour = TourServiceData(
            name="Обзорная экскурсия по Дубаю",
            name_en="Dubai City Tour",
            date="15.02.2025",
            time="09:00",
            duration="8 часов",
            pickup_location="Ваш отель в Дубае",
            participants=2,
            children=1,
            price_per_person=150,
            total_price=375,
            currency="AED",
            includes=[
                "Транспорт с кондиционером",
                "Профессиональный гид",
                "Вода в автомобиле",
                "Входные билеты (где указано)"
            ],
            excludes=[
                "Личные расходы",
                "Чаевые",
                "Обед"
            ]
        )

        return generator.create_tour_contract(
            client=client,
            tour_data=tour,
            prepayment_percent=50,
            cancellation_policy="standard",
            language=language
        )

    elif demo_type == "car":
        client = ClientData(
            full_name="Петров Петр Петрович",
            full_name_en="Petr Petrov",
            passport_number="76 9876543",
            phone="+7 999 987 65 43",
            email="petrov@example.com"
        )

        car = CarRentalData(
            vehicle_make="Mercedes-Benz",
            vehicle_model="S-Class",
            vehicle_year=2024,
            plate_number="A 12345",
            color="Черный",
            mileage_start=15000,
            rental_start="10.02.2025",
            rental_end="17.02.2025",
            pickup_location="Аэропорт DXB",
            dropoff_location="Аэропорт DXB",
            daily_rate=1500,
            total_days=7,
            total_price=10500,
            currency="AED",
            deposit_amount=5000,
            deposit_type="card_hold",
            insurance_type="full",
            insurance_included=True,
            excess_amount=2000,
            fuel_policy="full_to_full",
            special_conditions=[
                "Только для передвижения по ОАЭ",
                "Минимальный возраст водителя 25 лет"
            ]
        )

        return generator.create_car_rental_contract(
            client=client,
            car_data=car,
            language=language
        )

    elif demo_type == "yacht":
        client = ClientData(
            full_name="Сидоров Сидор Сидорович",
            full_name_en="Sidor Sidorov",
            passport_number="77 5555555",
            phone="+7 999 555 55 55",
            email="sidorov@example.com"
        )

        yacht = YachtCharterData(
            yacht_name="Ocean Dream",
            yacht_type="motor",
            yacht_length="85 футов",
            capacity=15,
            charter_date="20.02.2025",
            departure_time="14:00",
            return_time="18:00",
            duration_hours=4,
            departure_marina="Dubai Marina Yacht Club",
            route_description="Dubai Marina - Palm Jumeirah - Atlantis - Burj Al Arab",
            stops=["Palm Jumeirah (купание)", "Atlantis (фото)"],
            guests=10,
            crew_included=True,
            captain_name="Captain Ahmed",
            base_price=4000,
            per_hour_rate=1000,
            total_price=4000,
            currency="AED",
            deposit_amount=2000,
            catering_included=True,
            catering_description="BBQ на борту + напитки",
            catering_price=1500,
            fuel_included=True,
            special_requests=[
                "Украшение для дня рождения",
                "Торт с доставкой на борт"
            ]
        )

        return generator.create_yacht_charter_contract(
            client=client,
            yacht_data=yacht,
            prepayment_percent=50,
            language=language
        )

    elif demo_type == "agent":
        agent = AgentContractData(
            agency_name="Солнечный Путь",
            agency_name_en="Sunny Way Travel",
            trade_license="12345678",
            legal_address="Москва, ул. Тверская, д. 1",
            contact_person="Козлова Мария Ивановна",
            contact_position="Генеральный директор",
            contact_phone="+7 495 123 45 67",
            contact_email="info@sunnyway.ru",
            commission_rate=15,
            commission_type="percentage",
            payment_terms="prepaid",
            payment_period=7,
            contract_start="01.02.2025",
            contract_end="31.01.2026",
            auto_renewal=True,
            notice_period=30,
            services_covered=[
                "Экскурсии по ОАЭ",
                "Трансферы",
                "Чартер яхт",
                "Аренда автомобилей"
            ],
            exclusive=False,
            territory="Россия и СНГ",
            bank_name="Сбербанк",
            iban="RU123456789012345678901234567",
            swift="SABRRUMM"
        )

        return generator.create_agent_contract(
            agent_data=agent,
            language=language
        )


def output_contract(
    generator: ContractGenerator,
    contract: ContractData,
    format_type: str,
    include_qr: bool
):
    """Вывести/сохранить договор."""
    print(f"\n=== Договор создан ===")
    print(f"ID: {contract.contract_id}")
    print(f"Номер: {contract.contract_number}")
    print(f"Тип: {contract.contract_type.value}")
    print(f"Сумма: {contract.total_amount:,.2f} {contract.currency}")

    if format_type == "text":
        print("\n" + "=" * 60)
        print(generator.generate_text(contract))
        return

    try:
        if format_type in ("pdf", "both"):
            if REPORTLAB_AVAILABLE:
                pdf_path = generator.generate_pdf(contract, include_qr)
                print(f"\nPDF: {pdf_path}")
            else:
                print("\nPDF: reportlab не установлен (pip install reportlab)")

        if format_type in ("docx", "both"):
            if DOCX_AVAILABLE:
                docx_path = generator.generate_docx(contract, include_qr)
                print(f"DOCX: {docx_path}")
            else:
                print("DOCX: python-docx не установлен (pip install python-docx)")
    except Exception as e:
        print(f"\nОшибка генерации: {e}")
        print("Текстовая версия:")
        print(generator.generate_text(contract))


if __name__ == "__main__":
    main()
